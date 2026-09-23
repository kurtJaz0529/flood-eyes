# -*- coding: utf-8 -*-
"""填写附件7报名汇总表"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

SRC = r"C:\Users\21679\AppData\Local\Temp\notice_extract\地理与规划学院关于组织参加四川省国际大学生创新大赛（2026）的通知\附件7：地理与规划学院中国国际大学生创新大赛汇总表.xlsx"
DST = r"C:\Users\21679\Desktop\高教主赛道+慧眼识灾-遥感AI洪水智能识别系统+郝佳鑫.xlsx"

intro = ("针对洪灾应急人工解译遥感影像慢的痛点，本项目基于U-Net模型构建洪水淹没智能识别系统，"
         "用Sentinel-2影像实现分钟级自动提取与面积统计，经2020鄱阳湖特大洪水灾例验证，"
         "助力应急快速决策。")
print("简介字数：", len(intro))

row = [1,
       "慧眼识灾——基于深度学习的遥感洪水智能识别系统",
       "高教主赛道",
       "本科生创意组（“人工智能+”类）",
       "郝佳鑫",
       "遥感科学与技术",
       "2025级",
       "17711376089",
       "宋艺欣（2025级遥感一班）",
       "（待填）",
       "（待填）",
       intro]

wb = openpyxl.load_workbook(SRC)
ws = wb.active
for j, v in enumerate(row, start=1):
    ws.cell(row=2, column=j, value=v)
wb.save(DST)
print("已保存：", DST)

wb2 = openpyxl.load_workbook(DST)
print([c.value for c in wb2.active[2]])
